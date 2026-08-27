from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent: DocumentObject):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def clean(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split())


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    path = Path(sys.argv[1])
    start = int(sys.argv[2]) if len(sys.argv) > 2 else 1
    end = int(sys.argv[3]) if len(sys.argv) > 3 else 10**9
    doc = Document(path)

    print(f"DOCUMENT={path}")
    print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} INLINE_SHAPES={len(doc.inline_shapes)}")
    print("=== BODY ORDER ===")

    paragraph_index = 0
    table_index = 0
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            paragraph_index += 1
            text = clean(block.text)
            if not text:
                continue
            style = block.style.name if block.style is not None else ""
            if start <= paragraph_index <= end:
                print(f"[P{paragraph_index:04d}][{style}] {text}")
        else:
            table_index += 1
            if not (start <= paragraph_index <= end):
                continue
            print(f"[TABLE {table_index}] rows={len(block.rows)} cols={len(block.columns)}")
            for row in block.rows:
                values = [clean(cell.text) for cell in row.cells]
                print(" | ".join(values))
            print(f"[/TABLE {table_index}]")


if __name__ == "__main__":
    main()
